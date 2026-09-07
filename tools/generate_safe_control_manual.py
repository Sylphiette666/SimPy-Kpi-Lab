# ruff: noqa: E501

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "SimPy_KPI_Lab_安全闭环模型详细解释说明.docx"
QA_DIR = PROJECT_ROOT / ".docx_qa" / "safe_model_manual"
ARCHITECTURE_IMAGE = QA_DIR / "architecture.png"

DOCUMENT_VERSION = "0.2.0"
DOCUMENT_DATE = date(2026, 8, 28)
DESIGN_PRESET = "compact_reference_guide"
HEADER_PATTERN = "editorial_cover"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_CELL_TOP_BOTTOM_DXA = 80
TABLE_CELL_START_END_DXA = 120

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
TEAL = "2A7F83"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F3F7FB"
PALE_TEAL = "EAF5F5"
PALE_YELLOW = "FFF7D6"
PALE_RED = "FCE8E6"
LIGHT_GRAY = "F5F7FA"
MID_GRAY = "6B7280"
BORDER = "C7D2E0"
WHITE = "FFFFFF"
BLACK = "1F2937"

_BULLET_NUM_ID: int | None = None
_DECIMAL_NUM_ID: int | None = None


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
    ascii_font: str = "Calibri",
    east_asia_font: str = "Microsoft YaHei",
) -> None:
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{attribute}"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_language(paragraph, language: str = "zh-CN") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    lang = ppr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        ppr.append(lang)
    lang.set(qn("w:val"), language)
    lang.set(qn("w:eastAsia"), language)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(
    cell,
    *,
    top: int = TABLE_CELL_TOP_BOTTOM_DXA,
    start: int = TABLE_CELL_START_END_DXA,
    bottom: int = TABLE_CELL_TOP_BOTTOM_DXA,
    end: int = TABLE_CELL_START_END_DXA,
) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str = BORDER, size: int = 6) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "start", "bottom", "end"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths_dxa: Sequence[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tblpr = tbl.tblPr

    tblw = tblpr.first_child_found_in("w:tblW")
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")

    tblind = tblpr.first_child_found_in("w:tblInd")
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")

    layout = tblpr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = BLACK,
    size: float = 8.2,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.paragraph_format.widow_control = True
    set_paragraph_language(paragraph)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_caption(doc: Document, text: str) -> object:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    set_paragraph_language(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=8.5, italic=True, color=MID_GRAY)
    return paragraph


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    widths_dxa: Sequence[int],
    *,
    caption: str | None = None,
    font_size: float = 8.2,
) -> object:
    if caption:
        add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_header(header_row)
    for index, (cell, label) in enumerate(zip(header_row.cells, headers, strict=True)):
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_borders(cell)
        set_cell_margins(cell)
        set_cell_text(
            cell,
            label,
            bold=True,
            color=NAVY,
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT,
        )
    for row_data in rows:
        if len(row_data) != len(headers):
            raise ValueError(f"row length mismatch: {row_data}")
        row = table.add_row()
        prevent_row_split(row)
        for cell, value in zip(row.cells, row_data, strict=True):
            set_cell_borders(cell)
            set_cell_margins(cell)
            set_cell_text(cell, str(value), size=font_size)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer.paragraph_format.line_spacing = Pt(1)
    return table


def _next_numbering_id(numbering, tag: str, attribute: str) -> int:
    values = [int(item.get(qn(attribute))) for item in numbering.findall(qn(tag))]
    return max(values, default=0) + 1


def create_numbering_definition(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(numfmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "271")
    ppr.append(indent)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial" if bullet else "Calibri")
    rfonts.set(qn("w:hAnsi"), "Arial" if bullet else "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.append(rfonts)
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    direct_id = OxmlElement("w:numId")
    direct_id.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(direct_id)


def add_bullets(doc: Document, items: Sequence[str]) -> None:
    if _BULLET_NUM_ID is None:
        raise RuntimeError("numbering has not been initialized")
    for item in items:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, _BULLET_NUM_ID)
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.widow_control = True
        set_paragraph_language(paragraph)
        run = paragraph.add_run(item)
        set_run_font(run, size=11, color=BLACK)


def add_numbered(doc: Document, items: Sequence[str]) -> None:
    if _DECIMAL_NUM_ID is None:
        raise RuntimeError("numbering has not been initialized")
    for item in items:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, _DECIMAL_NUM_ID)
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.widow_control = True
        set_paragraph_language(paragraph)
        run = paragraph.add_run(item)
        set_run_font(run, size=11, color=BLACK)


def add_body(doc: Document, text: str, *, lead: str | None = None) -> object:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.widow_control = True
    set_paragraph_language(paragraph)
    if lead and text.startswith(lead):
        lead_run = paragraph.add_run(lead)
        set_run_font(lead_run, bold=True, color=NAVY)
        rest = paragraph.add_run(text[len(lead) :])
        set_run_font(rest, color=BLACK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, color=BLACK)
    return paragraph


def add_heading(doc: Document, text: str, level: int, *, page_break: bool = False) -> object:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.widow_control = True
    set_paragraph_language(paragraph)
    run = paragraph.add_run(text)
    set_run_font(
        run,
        size={1: 16, 2: 13, 3: 12}[level],
        bold=True,
        color=BLUE if level < 3 else DARK_BLUE,
    )
    return paragraph


def set_paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border(paragraph, *, color: str, side: str = "left", size: int = 18) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "6")
    border.set(qn("w:color"), color)


def add_callout(doc: Document, title: str, body: str, *, kind: str = "info") -> object:
    fill, accent = {
        "info": (PALE_BLUE, BLUE),
        "safe": (PALE_TEAL, TEAL),
        "warn": (PALE_YELLOW, "D69E2E"),
        "risk": (PALE_RED, "C0392B"),
    }[kind]
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.18
    paragraph.paragraph_format.keep_together = True
    set_paragraph_language(paragraph)
    set_paragraph_shading(paragraph, fill)
    set_paragraph_border(paragraph, color=accent)
    title_run = paragraph.add_run(f"{title}  ")
    set_run_font(title_run, bold=True, color=accent)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, color=BLACK)
    return paragraph


def add_code_block(doc: Document, text: str, *, font_size: float = 8.2) -> object:
    paragraph = doc.add_paragraph(style="Code Block")
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.widow_control = False
    set_paragraph_language(paragraph)
    set_paragraph_shading(paragraph, LIGHT_GRAY)
    set_paragraph_border(paragraph, color="D7DEE8", side="left", size=12)
    run = paragraph.add_run(text.rstrip())
    set_run_font(
        run,
        size=font_size,
        color="273444",
        ascii_font="Consolas",
        east_asia_font="Microsoft YaHei UI",
    )
    return paragraph


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    set_run_font(run, size=8, color=MID_GRAY)


def add_toc(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在 Word 中自动更新；如未更新，请按 Ctrl+A，再按 F9。"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])
    set_run_font(run, size=10, color=MID_GRAY)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.2)
    code.font.color.rgb = RGBColor.from_string("273444")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(5)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.08

    small = styles.add_style("Small Text", WD_STYLE_TYPE.PARAGRAPH)
    small.font.name = "Calibri"
    small.font.size = Pt(8.5)
    small.font.color.rgb = RGBColor.from_string(MID_GRAY)
    small._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.12


def configure_sections(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("SIMPY KPI LAB")
    set_run_font(left, size=8, bold=True, color=BLUE)
    right = paragraph.add_run("\t安全闭环模型技术说明")
    set_run_font(right, size=8, color=MID_GRAY)
    set_paragraph_border(paragraph, color=LIGHT_BLUE, side="bottom", size=5)

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    paragraph = first_footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(
        f"SimPy KPI Lab  •  v{DOCUMENT_VERSION}  •  {DOCUMENT_DATE.isoformat()}"
    )
    set_run_font(run, size=8, color=MID_GRAY)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run(f"技术说明 · v{DOCUMENT_VERSION}")
    set_run_font(left, size=8, color=MID_GRAY)
    page_label = paragraph.add_run("\t第 ")
    set_run_font(page_label, size=8, color=MID_GRAY)
    add_page_field(paragraph)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=8, color=MID_GRAY)


def set_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _font(size: int, *, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\msyhbd.ttc") if bold else Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                return ImageFont.truetype(
                    str(candidate), size=size, index=1 if bold and candidate.suffix == ".ttc" else 0
                )
            except OSError:
                return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_architecture_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1700, 980), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = _font(44, bold=True)
    box_title_font = _font(30, bold=True)
    box_font = _font(23)
    small_font = _font(20)

    draw.text((70, 42), "SimPy KPI Lab：可审计的下一轮实验闭环", font=title_font, fill="#17365D")
    draw.text(
        (72, 105),
        "AI 只提出结构化动作；本地策略、人工审批和版本校验决定是否生成下一版不可变配置。",
        font=small_font,
        fill="#536273",
    )

    boxes = [
        ((70, 205, 350, 390), "配置快照", "严格 Pydantic\n版本 + SHA-256"),
        ((430, 205, 710, 390), "实验运行", "参数网格\nSimPy replications"),
        ((790, 205, 1070, 390), "KPI 汇总", "均值 / 标准差\nSE / 置信区间"),
        ((1150, 205, 1630, 390), "OpenAI 提案（可选）", "Structured Outputs\n仅白名单动作数据"),
        ((1150, 535, 1630, 735), "人工审批", "approve / reject\nactor + reason + operation_id"),
        ((690, 535, 1070, 735), "本地安全应用", "allowlist + 严格校验\n工作量预算 + 哈希"),
        ((250, 535, 610, 735), "下一版配置", "单调 version\n只影响下一次运行"),
    ]

    for index, (rect, heading, detail) in enumerate(boxes):
        fill = "#F3F7FB"
        outline = "#2E74B5"
        if index == 4:
            fill, outline = "#FFF7D6", "#D69E2E"
        elif index == 5:
            fill, outline = "#EAF5F5", "#2A7F83"
        draw.rounded_rectangle(rect, radius=24, fill=fill, outline=outline, width=4)
        x1, y1, x2, _ = rect
        draw.text((x1 + 25, y1 + 25), heading, font=box_title_font, fill="#17365D")
        draw.multiline_text((x1 + 25, y1 + 88), detail, font=box_font, fill="#334155", spacing=8)

    def arrow(start: tuple[int, int], end: tuple[int, int], color: str = "#2E74B5") -> None:
        draw.line([start, end], fill=color, width=7)
        x2, y2 = end
        x1, y1 = start
        if abs(x2 - x1) >= abs(y2 - y1):
            direction = 1 if x2 > x1 else -1
            points = [(x2, y2), (x2 - direction * 22, y2 - 13), (x2 - direction * 22, y2 + 13)]
        else:
            direction = 1 if y2 > y1 else -1
            points = [(x2, y2), (x2 - 13, y2 - direction * 22), (x2 + 13, y2 - direction * 22)]
        draw.polygon(points, fill=color)

    arrow((350, 298), (430, 298))
    arrow((710, 298), (790, 298))
    arrow((1070, 298), (1150, 298))
    arrow((1390, 390), (1390, 535), "#D69E2E")
    arrow((1150, 635), (1070, 635), "#2A7F83")
    arrow((690, 635), (610, 635), "#2A7F83")
    draw.line([(250, 635), (125, 635), (125, 390)], fill="#2A7F83", width=7)
    arrow((125, 390), (125, 390), "#2A7F83")
    draw.text(
        (93, 760), "REST + WebSocket：管理会话、运行、审批、审计", font=small_font, fill="#536273"
    )
    draw.text(
        (930, 810),
        "MCP：读取 / 校验 / 提交待审批提案（不暴露运行与批准）",
        font=small_font,
        fill="#536273",
    )
    draw.rounded_rectangle((70, 900, 1630, 940), radius=12, fill="#17365D")
    draw.text(
        (360, 903),
        "核心边界：不执行模型生成的代码，不热改正在运行的 SimPy generator",
        font=small_font,
        fill="#FFFFFF",
    )
    image.save(path, dpi=(180, 180))


def add_architecture_figure(doc: Document) -> None:
    draw_architecture_diagram(ARCHITECTURE_IMAGE)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    shape = run.add_picture(str(ARCHITECTURE_IMAGE), width=Inches(6.35))
    shape._inline.docPr.set(
        "descr",
        "安全闭环架构图：不可变配置进入 SimPy 多次实验，KPI 汇总供 OpenAI 生成结构化白名单动作，人工审批后经本地策略、工作量、哈希和版本校验形成下一版配置；REST/WebSocket 管理流程，MCP 仅提供受限提案接口。",
    )
    shape._inline.docPr.set("title", "SimPy KPI Lab 安全闭环架构")
    add_caption(doc, "图 1　当前模型的安全闭环：AI 负责建议，服务端与人工共同决定下一轮配置。")


def add_cover(doc: Document) -> None:
    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(56)
    eyebrow.paragraph_format.space_after = Pt(18)
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("SIMPY KPI LAB  ·  TECHNICAL MANUAL")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    set_paragraph_language(title)
    run = title.add_run("安全闭环仿真模型\n详细解释说明")
    set_run_font(run, size=30, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle.paragraph_format.line_spacing = 1.2
    set_paragraph_language(subtitle)
    run = subtitle.add_run(
        "SimPy + KPI 统计 + 多次实验 + OpenAI 结构化提案\n+ 人工审批 + REST / WebSocket + 受限 MCP"
    )
    set_run_font(run, size=14, color=TEAL)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(18)
    set_paragraph_border(rule, color=BLUE, side="bottom", size=18)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Inches(0.55)
    summary.paragraph_format.right_indent = Inches(0.55)
    summary.paragraph_format.space_after = Pt(22)
    summary.paragraph_format.line_spacing = 1.35
    set_paragraph_language(summary)
    run = summary.add_run(
        "面向开发、实验设计、运营分析与审批人员的实现级说明。文档依据当前仓库 v0.2.0 代码编写，重点解释可复现实验、KPI 口径、AI 安全边界、人工审批状态机以及本地部署方法。"
    )
    set_run_font(run, size=11.5, color=BLACK)

    add_table(
        doc,
        ["项目项", "当前值"],
        [
            ("软件版本", DOCUMENT_VERSION),
            ("说明日期", DOCUMENT_DATE.isoformat()),
            ("适用模型", "安全的下一轮实验控制，而非运行中热修改"),
            ("设计预设", f"{DESIGN_PRESET} / {HEADER_PATTERN}"),
        ],
        [2300, 7060],
        font_size=9,
    )
    add_callout(
        doc,
        "一句话边界",
        "OpenAI 只能返回经过 Schema 约束的建议数据；任何配置变化都必须经过人工决定与本地校验，并且只作用于下一次仿真。",
        kind="safe",
    )


def add_reader_guide(doc: Document) -> None:
    add_heading(doc, "文档说明与阅读地图", 1, page_break=True)
    add_body(
        doc,
        "本说明书从“模型为何这样设计”开始，逐层下钻到配置、仿真、实验、KPI、OpenAI、审批、API、MCP 与运维。它既可作为新成员入门材料，也可作为接口联调和安全评审的参考。",
    )
    add_heading(doc, "一分钟理解", 2)
    add_numbered(
        doc,
        [
            "用 YAML 或 JSON 定义一个不可变 ProjectConfig 快照。",
            "把参数网格展开为多个场景；每个场景执行多次 replication。",
            "SimPy 运行串行多工位排队模型，本地代码记录事件并计算 KPI。",
            "统计层按场景汇总均值、样本标准差、标准误与置信区间。",
            "可选地把当前配置、KPI、目标与精确动作白名单交给 OpenAI，得到结构化提案。",
            "人工逐项批准或拒绝；批准后仍需通过 allowlist、全配置、工作量、哈希和版本校验。",
            "形成下一版配置，再开始下一轮实验；已经运行的 replication 不会被热修改。",
        ],
    )
    add_heading(doc, "按角色阅读", 2)
    add_table(
        doc,
        ["角色", "优先章节", "要回答的问题"],
        [
            ("仿真建模人员", "配置、SimPy、扩展指南", "实体如何流动？随机性如何隔离？"),
            ("数据/运营分析人员", "实验设计、KPI、结果解读", "指标口径和不确定性是否可信？"),
            ("AI / API 开发人员", "OpenAI、REST、MCP", "输入输出契约和错误边界是什么？"),
            ("审批与安全人员", "白名单、状态机、审计", "AI 能做什么、不能做什么？"),
            ("运维人员", "PowerShell、故障排查、限制", "如何本地启动、验证并定位问题？"),
        ],
        [1650, 3000, 4710],
        caption="表 1　不同角色的建议阅读路径",
    )
    add_callout(
        doc,
        "适用性说明",
        "当前业务模型是“所有实体依次经过全部工位”的串行网络。路由、返工、优先级、批处理、故障维修等能力需要在此框架上扩展，不应被误认为已经内置。",
        kind="warn",
    )

    add_heading(doc, "目录", 1, page_break=True)
    add_toc(doc)
    add_body(
        doc,
        "提示：本文件包含自动目录字段。Word 通常会在打开时更新；若页码尚未刷新，请按 Ctrl+A 选中全文，再按 F9。",
    )


def add_positioning(doc: Document) -> None:
    add_heading(doc, "模型定位与设计原则", 1, page_break=True)
    add_body(
        doc,
        "SimPy KPI Lab 是一个可复现的离散事件仿真实验框架，并在其上增加可选的 OpenAI 决策辅助与人工审批控制面。它把“生成建议”和“执行变化”分开：模型可以解释证据、提出有限动作，但不能执行 Python、调用 shell、直接改文件或绕过审批运行新方案。",
    )
    add_heading(doc, "要解决的核心问题", 2)
    add_bullets(
        doc,
        [
            "KPI 可统计：每个 replication 输出一致口径的业务与工位指标，并给出跨重复实验的不确定性。",
            "实验可重复：固定基准 seed，使用稳定命名空间派生随机流，避免 Python 进程哈希随机化影响结果。",
            "方案可比较：参数网格自动展开；共同随机数可让不同方案共享对应随机输入。",
            "AI 可接入：Responses API 使用 Pydantic Structured Outputs 返回可解析的分析或动作提案。",
            "控制可审计：人工审批、乐观版本、幂等 operation_id、配置哈希与顺序事件记录共同构成审计链。",
        ],
    )
    add_callout(
        doc,
        "不是自治执行器",
        "系统没有让 OpenAI 生成并运行 Python 的路径，也没有“模型调用 approve”能力。它是一个带安全门的实验协作系统。",
        kind="safe",
    )
    add_callout(
        doc,
        "正确的安全边界",
        "应通过 SimulationControlService + policy.py 管理配置变化。单独调用公开的 ApprovalWorkflow 只保证动作是严格、惰性的结构化数据，并不会自行执行完整 ProjectConfig allowlist 或工作量校验。",
        kind="warn",
    )
    add_heading(doc, "关键设计原则", 2)
    add_table(
        doc,
        ["原则", "落地方式", "直接收益"],
        [
            ("仿真与 AI 解耦", "没有 API key 仍可运行全部仿真与 KPI", "离线可用，成本和故障隔离"),
            ("配置不可变快照", "批准只生成下一版本，不触碰在途 generator", "可复现、可追责"),
            ("结构化而非代码", "AI 输出 ActionProposal 数据模型", "消除任意代码执行入口"),
            (
                "双重校验",
                "Schema 解析后再由本地 allowlist 与 ProjectConfig 校验",
                "防止格式正确但语义越权",
            ),
            ("显式不确定性", "n、n_missing、std、SE、CI 均输出", "避免把一次随机结果当结论"),
            ("最小权限接口", "MCP 不暴露运行、批准、拒绝和建会话", "降低智能客户端越权面"),
        ],
        [1650, 4250, 3460],
        caption="表 2　从设计原则到实现机制",
    )


def add_architecture(doc: Document) -> None:
    add_heading(doc, "总体架构与端到端闭环", 1)
    add_architecture_figure(doc)
    add_heading(doc, "分层说明", 2)
    add_table(
        doc,
        ["层", "主要模块", "职责", "不负责"],
        [
            (
                "配置层",
                "config.py",
                "解析分布、工位、实验与 OpenAI 参数；拒绝未知字段",
                "不运行仿真",
            ),
            (
                "仿真层",
                "simulation.py / rng.py",
                "SimPy 实体流程、资源竞争、独立随机流",
                "不调用 AI",
            ),
            ("指标层", "kpi.py", "事件采集、窗口裁剪、右删失与 KPI", "不修改 KPI 值"),
            ("实验层", "experiment.py", "场景展开、replication、并行、聚合与导出", "不审批动作"),
            ("AI 适配层", "ai.py / agent.py", "结构化分析与白名单提案", "不执行动作"),
            (
                "安全策略层",
                "policy.py",
                "allowlist、配置哈希、工作量预算、应用配置副本",
                "不绕过人工",
            ),
            (
                "工作流层",
                "workflow.py / service.py",
                "状态机、版本、幂等、审计、服务会话",
                "不热改在途运行",
            ),
            ("接口层", "api.py / mcp_server.py", "REST、WebSocket 与受限 MCP", "不扩大底层权限"),
        ],
        [1300, 2140, 3340, 2580],
        caption="表 3　架构分层及职责边界",
        font_size=7.7,
    )
    add_heading(doc, "一次完整闭环", 2)
    add_numbered(
        doc,
        [
            "服务创建会话，深拷贝配置并记录 version 与 canonical SHA-256 config_hash。",
            "调用 runs 接口时，以当前配置快照构造 ExperimentRunner；服务端分配隔离输出目录。",
            "实验结束后保存 results.json、replications.csv、summary.csv，并把最新 KPI 摘要留在会话中。",
            "人工提交提案，或请求 OpenAI 根据当前配置、最新 KPI、目标与 allowed_actions 生成计划。",
            "每个动作成为 pending 提案；操作者按顺序 approve 或 reject。",
            "批准操作先检查 expected_version、提案创建时哈希、动作顺序与本地政策，再把动作应用到配置副本。",
            "成功后 proposal 进入 applied，配置 version 单调增加；失败则进入 failed 并留下错误与审计事件。",
            "下一次 run 读取新快照；此前的结果、seed、配置与审计记录仍可追溯。",
        ],
    )
    add_callout(
        doc,
        "运行时边界",
        "这里的“动态控制”指版本化的下一轮实验控制，不是暂停、序列化或修改正在运行的 SimPy 进程。该选择牺牲即时热调参，换取确定性、可复现性和安全性。",
        kind="info",
    )


def add_configuration(doc: Document) -> None:
    add_heading(doc, "配置模型与数据契约", 1, page_break=True)
    add_body(
        doc,
        '所有配置模型均使用 Pydantic 且设置 extra="forbid"。未知字段会立即报错，避免拼写错误悄悄变成无效参数。初始 YAML/JSON 载入使用 Pydantic 的常规解析，允许部分安全的类型转换；批准动作应用到配置副本时才使用 strict=True。配置既可从 YAML 读取，也可通过 REST 以 JSON 提交；ProjectConfig 是所有执行与审批的边界对象。',
    )
    add_heading(doc, "ProjectConfig 顶层", 2)
    add_table(
        doc,
        ["字段", "类型 / 默认", "含义", "关键约束"],
        [
            ("project_name", "str / simpy-kpi-lab", "项目显示名称", "未知字段被拒绝"),
            ("simulation", "SimulationConfig / 必填", "单个场景的运行模型", "必须至少一个工位"),
            ("experiment", "ExperimentConfig / 默认", "replication 与场景设计", "网格项不可为空"),
            ("openai", "OpenAIConfig / 默认", "可选 AI 调用参数", "key 不属于配置"),
        ],
        [1900, 2500, 2860, 2100],
        caption="表 4　ProjectConfig 顶层字段",
    )
    add_heading(doc, "SimulationConfig", 2)
    add_table(
        doc,
        ["字段", "默认 / 范围", "说明"],
        [
            ("name", "service_system", "仿真模型名称"),
            ("until", "> 0，必填", "仿真终止时间"),
            ("warmup", "0；0 ≤ warmup < until", "KPI 统计预热期"),
            ("first_arrival_at_zero", "true", "首个实体是否在 t=0 到达"),
            ("max_arrivals", "null 或 ≥ 1", "每次 replication 的到达上限"),
            ("arrival_interarrival", "DistributionConfig", "相邻到达间隔分布，必须严格正支撑"),
            ("stations", "非空列表", "实体按列表顺序访问所有工位；名称唯一"),
            ("cycle_time_target", "null 或 > 0", "service_level 的周期时间目标"),
        ],
        [2200, 2600, 4560],
        caption="表 5　仿真配置字段",
    )
    add_heading(doc, "ExperimentConfig 与 OpenAIConfig", 2)
    add_table(
        doc,
        ["分组", "字段", "默认", "说明"],
        [
            ("experiment", "replications", "10", "每个场景的独立重复次数"),
            ("experiment", "base_seed", "20260825", "稳定随机派生的根 seed"),
            (
                "experiment",
                "common_random_numbers",
                "true",
                "不同场景同编号 replication 是否共享随机输入",
            ),
            ("experiment", "confidence_level", "0.95", "0 到 1 之间的聚合置信水平"),
            ("experiment", "parameter_grid", "{}", "相对 simulation 的点路径到候选值列表"),
            ("experiment", "output_dir", "outputs", "CLI 输出目录；控制 API 会忽略请求值"),
            (
                "openai",
                "model",
                "gpt-5.6",
                "Responses API 模型名；仅 ProjectConfig.load 路径可被 SIMLAB_OPENAI_MODEL 覆盖",
            ),
            ("openai", "max_output_tokens", "2500", "结构化响应输出上限，至少 256"),
            ("openai", "timeout_seconds", "60", "客户端请求超时"),
            ("openai", "max_retries", "2", "SDK 重试次数，0–10"),
            ("openai", "store", "false", "是否允许服务端存储响应；默认关闭"),
        ],
        [1400, 2200, 1600, 4160],
        caption="表 6　实验与 OpenAI 配置",
        font_size=7.8,
    )
    add_heading(doc, "分布配置", 2)
    add_table(
        doc,
        ["kind", "必需参数", "有效条件", "采样方式"],
        [
            ("exponential", "mean", "mean > 0", "expovariate(1 / mean)"),
            ("deterministic", "value", "value ≥ 0；到达间隔须 > 0", "固定值"),
            ("uniform", "low, high", "0 ≤ low ≤ high；到达 low 不得为 0", "uniform(low, high)"),
            (
                "triangular",
                "low, mode, high",
                "0 ≤ low ≤ mode ≤ high；到达 low 不得为 0",
                "triangular(low, high, mode)",
            ),
        ],
        [1800, 2100, 3400, 2060],
        caption="表 7　内置随机分布",
    )
    add_callout(
        doc,
        "配置整洁性",
        "模型会拒绝未知字段，但不会禁止同时填写与当前 kind 无关的已知分布字段。例如 exponential 可额外带 low；该值不参与采样，却可能进入动态动作目录。建议只填写当前分布真正使用的参数。",
        kind="warn",
    )
    add_heading(doc, "控制 API 示例配置", 2)
    add_code_block(
        doc,
        r"""{
  "project_name": "api-service-center-demo",
  "simulation": {
    "name": "two-stage-service-center",
    "until": 120,
    "warmup": 20,
    "first_arrival_at_zero": true,
    "max_arrivals": 100,
    "arrival_interarrival": {"kind": "exponential", "mean": 5.0},
    "stations": [
      {"name": "registration", "capacity": 1,
       "service_time": {"kind": "exponential", "mean": 3.0}},
      {"name": "specialist", "capacity": 2,
       "service_time": {"kind": "triangular", "low": 4.0, "mode": 8.0, "high": 15.0}}
    ],
    "cycle_time_target": 30
  },
  "experiment": {"replications": 5, "base_seed": 20260825,
    "common_random_numbers": true, "confidence_level": 0.95,
    "parameter_grid": {}, "output_dir": "ignored-by-control-api"},
  "openai": {"model": "gpt-5.6", "max_output_tokens": 1500,
    "timeout_seconds": 60, "max_retries": 2, "store": false}
}""",
        font_size=7.3,
    )


def add_simulation(doc: Document) -> None:
    add_heading(doc, "SimPy 运行机制", 1)
    add_body(
        doc,
        "每个 replication 创建独立的 simpy.Environment、工位 Resource、KPICollector 与随机流。实体由 arrivals() 过程产生，每个 customer() 按 stations 列表依次请求资源、等待、接受服务，最后完成。env.run(until=simulation.until) 给出明确终止边界。",
    )
    add_heading(doc, "事件时序", 2)
    add_numbered(
        doc,
        [
            "到达生成器决定首个到达时点，并在每次到达后抽取下一间隔。",
            "customer 到达时记录 arrival；随后逐站记录 queue_enter。",
            "获得 Resource 后记录 service_start，由此计算该次等待与队长面积。",
            "从工位专属随机流抽取服务时长；记录与统计窗口重叠的 busy interval。",
            "完成所有工位后记录 completion、周期时间和跨工位总等待。",
            "仿真终止时 finalize 把仍在队列与仍在系统中的实体纳入 WIP、队长面积与删失指标。",
        ],
    )
    add_heading(doc, "命名随机流", 2)
    add_body(
        doc,
        "随机 seed 不是用 Python 内置 hash() 派生，而是由 rng.derive_seed 使用 BLAKE2b 对“base seed + namespace”稳定派生。一个 replication 先得到 replication seed，再分别派生 arrivals 与 service:<station_name> 随机流。",
    )
    add_table(
        doc,
        ["层级", "命名空间示例", "目的"],
        [
            ("replication", "replication:3", "同一编号重复实验可稳定重放"),
            ("非 CRN 场景", "replication:3|scenario:{...}", "场景也参与 seed，随机输入不配对"),
            ("到达流", "arrivals", "容量或工位变化不改变到达序列"),
            ("工位服务流", "service:specialist", "其他工位变化不扰动该工位随机序列"),
        ],
        [1900, 3700, 3760],
        caption="表 8　随机流的命名空间层级",
    )
    add_callout(
        doc,
        "复现条件",
        "要复现某次结果，需要同时保留完整配置、场景参数、replication 编号、实际 seed、软件版本与随机流方法。results.json 已记录其中的主要部分。",
        kind="info",
    )
    add_heading(doc, "终止与边界行为", 2)
    add_bullets(
        doc,
        [
            "到达时间满足 env.now < until；若下一间隔使到达时刻达到或超过 until，则不再创建实体。",
            "max_arrivals 是硬上限，可避免极短到达间隔导致单次实验实体数量失控。",
            "到达分布必须具有严格正的间隔支撑，防止零间隔无限循环。",
            "未在 until 前完成的实体不会获得完整周期时间，但会进入 WIP 与删失统计。",
        ],
    )


def add_experiments(doc: Document) -> None:
    add_heading(doc, "多次实验、参数网格与统计聚合", 1)
    add_heading(doc, "场景展开", 2)
    add_body(
        doc,
        "parameter_grid 的每个键是相对 simulation 的点路径，列表下标以十进制数字表示。所有候选值做笛卡尔积；例如 2 个容量 × 3 个到达均值会生成 6 个场景。每个场景再执行 replications 次，总任务数等于场景数 × replications。",
    )
    add_code_block(
        doc,
        "parameter_grid:\n  stations.0.capacity: [1, 2]\n  arrival_interarrival.mean: [4.0, 5.0, 6.0]",
    )
    add_heading(doc, "共同随机数（CRN）", 2)
    add_table(
        doc,
        ["设置", "seed 关系", "适用情形", "注意事项"],
        [
            (
                "true（默认）",
                "不同场景同编号 replication 共享 seed",
                "容量、服务时间等相近方案对比",
                "适合配对比较；仍应看差值分布",
            ),
            (
                "false",
                "场景参数指纹参与 seed",
                "场景结构差异大或不希望配对",
                "方案差异估计方差通常更大",
            ),
        ],
        [1700, 3000, 2780, 1880],
        caption="表 9　共同随机数策略",
    )
    add_body(
        doc,
        "当前 summary.csv 仍按场景分别计算均值与区间，不自动输出配对差值或配对置信区间。需要完整利用 CRN 时，应按 replication 编号连接 replications.csv，再分析方案差值。",
    )
    add_heading(doc, "并行执行", 2)
    add_body(
        doc,
        "workers=1 时当前进程顺序运行；workers>1 时使用 ProcessPoolExecutor。每个任务携带可序列化的 simulation、seed、replication、scenario 与 parameters，因此进程之间不共享 SimPy 环境或 Resource。Windows 上建议从 simlab CLI 启动，而不是在交互式解释器中直接创建进程池。",
    )
    add_heading(doc, "聚合公式", 2)
    add_body(doc, "对某个“场景 × KPI”，仅使用非缺失数值 x₁…xₙ：")
    add_code_block(
        doc,
        "mean = Σxᵢ / n\ns = sqrt[ Σ(xᵢ − mean)² / (n − 1) ]\nSE = s / sqrt(n)\nCI = mean ± z((1 + confidence_level) / 2) × SE",
        font_size=9,
    )
    add_bullets(
        doc,
        [
            "标准差使用样本标准差；仅有一个有效 replication 时 std、SE 与 CI 为 null。",
            "n 是有效样本数，n_total 是该场景全部 replication 数，n_missing = n_total − n。",
            "ci_method 固定记录为 normal_approximation；replication 很少或分布重尾时，应考虑增加次数或改用 t / bootstrap。",
            "min 与 max 用于快速查看极端重复实验，不应替代分位数或置信区间。",
            "比例指标的正态近似区间不会自动裁剪到 [0,1]；越界时应更换比例区间方法。",
            "n_missing 表示该 KPI 在 replication 中为 null，不表示系统保留了失败任务；任一任务抛错会使整个 runner 失败。",
        ],
    )


def add_kpis(doc: Document) -> None:
    add_heading(doc, "KPI 体系与统计口径", 1)
    add_callout(
        doc,
        "统一窗口",
        "所有窗口型指标以 [warmup, until) 为统计区间，窗口长度为 until − warmup。时间单位由业务统一定义；示例按“分钟”理解。",
        kind="info",
    )
    add_heading(doc, "系统级 KPI", 2)
    add_table(
        doc,
        ["指标", "角色 / 方向", "单位", "精确定义"],
        [
            ("arrivals", "context / 仅上下文", "count", "窗口内到达实体数"),
            ("completed", "context / 仅上下文", "count", "窗口内完成全部工位的实体数"),
            (
                "throughput_per_time_unit",
                "primary / 越高越好",
                "count / time",
                "窗口内完成数 ÷ 窗口长度",
            ),
            ("wip_end", "guardrail / 越低越好", "count", "until 时仍在系统内的实体数"),
            ("avg_cycle_time", "driver / 越低越好", "time", "有效完成 cohort 的平均总周期时间"),
            ("p50_cycle_time", "driver / 越低越好", "time", "有效完成 cohort 周期时间 P50"),
            ("p95_cycle_time", "primary / 越低越好", "time", "有效完成 cohort 周期时间 P95"),
            ("avg_wait_time", "driver / 越低越好", "time", "所有有效工位访问的平均单次等待"),
            (
                "avg_total_wait_time",
                "driver / 越低越好",
                "time",
                "有效完成实体跨全部工位的平均总等待",
            ),
            (
                "service_level",
                "primary / 越高越好",
                "ratio",
                "完成 cohort 中周期时间 ≤ target 的比例；无 target 时缺失",
            ),
            (
                "cycle_completion_fraction",
                "data_quality / 越高越好",
                "ratio",
                "窗口内到达 cohort 在 until 前完成的比例",
            ),
            (
                "censored_cycle_count",
                "data_quality / 越低越好",
                "count",
                "窗口内到达但直到 until 仍未完成的实体数",
            ),
        ],
        [2500, 2220, 1340, 3300],
        caption="表 10　系统级 KPI 目录",
        font_size=7.45,
    )
    add_heading(doc, "工位级 KPI", 2)
    add_table(
        doc,
        ["路径", "角色 / 方向", "公式或口径"],
        [
            (
                "station.<name>.utilization",
                "driver / 目标区间",
                "窗口内忙碌资源时间 ÷（capacity × 窗口长度）",
            ),
            (
                "station.<name>.avg_queue_length",
                "driver / 越低越好",
                "排队时间面积（含期末仍在排队者）÷ 窗口长度",
            ),
            (
                "station.<name>.avg_wait_time",
                "driver / 越低越好",
                "warmup 后入队且 until 前开始服务的访问平均等待",
            ),
            (
                "station.<name>.p95_wait_time",
                "guardrail / 越低越好",
                "同一有效访问集合的等待时间 P95",
            ),
        ],
        [3200, 2400, 3760],
        caption="表 11　每个工位重复生成的 KPI",
    )
    add_heading(doc, "cohort 与右删失", 2)
    add_body(
        doc,
        "周期 cohort 定义为“到达时刻在 warmup 及之后、且早于 until 的实体”。只有在 until 前完成者才有可观测的完整周期时间，因此 avg/p50/p95_cycle_time 与 service_level 不包含右删失实体。系统同时报告 censored_cycle_count 和 cycle_completion_fraction，防止尾部积压被悄悄忽略。",
    )
    add_callout(
        doc,
        "解读原则",
        "如果 cycle_completion_fraction 较低，即使已完成实体的 P95 看起来很好，也不能直接断言系统尾部体验良好。应延长 until、减少 warmup、提高容量或采用删失数据方法后再比较。",
        kind="warn",
    )
    add_heading(doc, "百分位数与空样本", 2)
    add_bullets(
        doc,
        [
            "P50/P95 使用排序后的线性插值：index=(n−1)×p。",
            "没有有效样本时返回 null，而不是 0；0 会错误暗示“没有等待”。",
            "工位 utilization 始终可算，但空工位的等待指标可能为 null。",
            "service_starts、observed_cycle_count 和 capacity 属于 replication 元数据，不进入跨实验 KPI catalog 汇总；当前 service_starts 实际是有效等待样本数，并非所有物理开工事件。",
        ],
    )


def add_openai(doc: Document) -> None:
    add_heading(doc, "OpenAI 接入：分析与结构化提案", 1)
    add_body(
        doc,
        "项目提供两个相互独立的 OpenAI 用途：一是对已有 KPI 结果生成结构化解释与 Markdown；二是在控制会话中生成待人工审批的动作计划。两者都使用 Responses API 与 Pydantic 结构化输出，本地 KPI 数值始终由 Python 计算。",
    )
    add_heading(doc, "两类适配器", 2)
    add_table(
        doc,
        ["用途", "实现", "输入", "输出", "是否改变配置"],
        [
            (
                "KPI 分析",
                "ai.py / OpenAIKPIAnalyst",
                "results.json、问题",
                "KPIAnalysis JSON + Markdown",
                "否",
            ),
            (
                "动作提案",
                "agent.py / OpenAIProposalAgent",
                "当前配置、KPI、目标、白名单",
                "ActionProposal",
                "否；仅生成 pending",
            ),
        ],
        [1500, 2350, 2380, 1850, 1280],
        caption="表 12　OpenAI 的两个使用面",
        font_size=7.8,
    )
    add_heading(doc, "Structured Outputs 调用", 2)
    add_code_block(
        doc,
        "response = client.responses.parse(\n    model=config.openai.model,\n    input=[system_message, user_payload],\n    text_format=ActionProposal,\n    max_output_tokens=config.openai.max_output_tokens,\n    store=config.openai.store,\n)\nproposal = response.output_parsed",
    )
    add_body(
        doc,
        "提案请求只序列化 current_config、kpi_summary、objective 和 allowed_actions。系统提示明确要求：动作类型仅限 set_parameter、change_policy、enable_resource；target 必须与白名单逐字一致；requires_approval 必须为 true；证据不足时返回空 actions。",
    )
    add_heading(doc, "ActionProposal 数据契约", 2)
    add_table(
        doc,
        ["对象", "字段", "限制"],
        [
            ("ActionProposal", "summary", "中文摘要字符串"),
            ("ActionProposal", "actions", "最多 10 项 ProposedAction"),
            ("ActionProposal", "caveats", "证据不足、风险和假设"),
            ("ProposedAction", "action_type", "三类动作之一"),
            ("ProposedAction", "target", "必须命中当前 allowlist"),
            ("ProposedAction", "proposed_value", "str / int / float / bool / null"),
            ("ProposedAction", "rationale / expected_effect", "理由与预期效果"),
            ("ProposedAction", "risk", "low / medium / high"),
            ("ProposedAction", "requires_approval", "必须为 true"),
        ],
        [2100, 3000, 4260],
        caption="表 13　结构化提案字段",
    )
    add_heading(doc, "密钥与客户端设置", 2)
    add_code_block(
        doc,
        '$env:OPENAI_API_KEY = "sk-请替换为真实密钥"\n.\\.venv\\Scripts\\simlab.exe analyze .\\outputs\\service_center\\results.json --model "gpt-5.6"',
    )
    add_callout(
        doc,
        "PowerShell 易错点",
        "环境变量名必须写成 $env:OPENAI_API_KEY。不要写 OPENAI\\_API\\_KEY；反斜杠是 Markdown 转义残留，不是变量名的一部分。不要把真实 key 写入 YAML、.env.example、提交记录或截图。",
        kind="risk",
    )
    add_bullets(
        doc,
        [
            "客户端默认超时 60 秒、最多重试 2 次；这些值可在 openai 配置中调整。",
            "控制 API 创建提案 Agent 时会把 max_output_tokens 再截到最多 4,000；直接构造 Agent 的类默认值是 1,500。",
            "SIMLAB_OPENAI_MODEL 仅影响通过 ProjectConfig.load 读取本地配置的路径；REST 提交的配置和独立 analyze 命令不把它当全局覆盖。",
            "默认 gpt-5.6 只是项目配置值，实际可用性取决于当前 API 项目权限与官方模型支持。",
            "store 默认为 false；若明确改为 true，应重新评估数据保留要求。",
            "认证、限流、超时、连接、HTTP 状态、拒答或结构解析失败均转换为可读错误；不会自动应用任何动作。",
        ],
    )


def add_policy(doc: Document) -> None:
    add_heading(doc, "动作白名单与本地安全策略", 1, page_break=True)
    add_heading(doc, "允许动作", 2)
    add_table(
        doc,
        ["动作", "目标形式", "value", "本地效果"],
        [
            ("set_parameter", "精确点路径", "JSON 标量或 null", "在配置副本上设置单一参数"),
            ("change_policy", "精确政策路径", "当前实现要求 bool", "切换 CRN 或首达时点政策"),
            (
                "enable_resource",
                "已登记 station.name",
                "null / bool / 正整数",
                "启停工位；正整数同时设容量",
            ),
        ],
        [1800, 2600, 2240, 2720],
        caption="表 14　仅有的三类动作",
    )
    add_heading(doc, "当前精确 target 目录", 2)
    add_table(
        doc,
        ["类别", "允许 target"],
        [
            (
                "仿真边界",
                "simulation.until；simulation.warmup；simulation.max_arrivals；simulation.cycle_time_target",
            ),
            ("实验统计", "experiment.replications；experiment.confidence_level"),
            ("布尔政策", "experiment.common_random_numbers；simulation.first_arrival_at_zero"),
            ("到达分布", "仅当前分布中非 null 的 mean / value / low / high / mode"),
            ("工位参数", "simulation.stations.<index>.capacity 与当前 service_time 的非 null 参数"),
            ("资源启停", "每个已登记 station.name"),
        ],
        [2100, 7260],
        caption="表 15　allowlist 的动态生成规则",
    )
    add_callout(
        doc,
        "明确禁止",
        "模型不能修改 OPENAI_API_KEY、openai.model、experiment.output_dir、experiment.base_seed、parameter_grid 或任意未登记点路径；没有 shell、文件、Python、eval、exec 或 subprocess 执行工具。",
        kind="risk",
    )
    add_heading(doc, "应用动作的检查顺序", 2)
    add_numbered(
        doc,
        [
            "验证动作是严格、有限、可序列化的判别联合；拒绝 NaN、Infinity、dunder 路径与未知字段。",
            "确认 action_type + target 精确命中当前配置生成的 allowlist。",
            "把动作应用到深拷贝后的配置数据；原配置在全部检查通过前不变。",
            "用 strict=True 重新构造完整 ProjectConfig，执行分布、warmup、工位名称等全部约束。",
            "对基础配置和参数网格展开后的每个场景执行工作量验证。",
            "只有全部成功才替换会话的当前配置，并生成新的 SHA-256 配置哈希。",
        ],
    )
    add_heading(doc, "有序多动作计划", 2)
    add_body(
        doc,
        "OpenAI 可以返回多项 actions，但服务端把每项动作转换成独立 proposal，并按顺序审批。后一步绑定前一步预览应用后的 config_hash。它不是“一次批准全部”的原子事务；若会改变配置的前一步被拒绝、失败或被其他操作取代，后续提案会因哈希过期而拒绝。若前一步本来就是不改变配置的 no-op，哈希相同，后一步不一定仅凭哈希失效。",
    )
    add_callout(
        doc,
        "工位下标规则",
        "enable_resource 必须排在所有 simulation.stations.<index> 的 set_parameter 之后。停用工位会改变列表下标，服务端通过顺序检查与 parameter_grid 重映射避免动作误落到其他工位。",
        kind="warn",
    )


def add_workload(doc: Document) -> None:
    add_heading(doc, "工作量护栏与资源预算", 1)
    add_body(
        doc,
        "控制 API 是长寿命服务，不能只依赖 Pydantic 的字段范围。PolicyLimits 进一步限制总任务与估算事件数，并且对 parameter_grid 展开的每个场景单独检查，防止网格值绕过基础配置限制。",
    )
    add_table(
        doc,
        ["限制", "默认值", "保护对象"],
        [
            ("max_replications", "500", "单场景 replication 数"),
            ("max_scenarios", "64", "参数网格笛卡尔积场景数"),
            ("max_total_replications", "2,000", "场景数 × replications"),
            ("max_until", "1,000,000", "每个展开场景的仿真时长"),
            ("max_stations", "50", "每个展开场景的工位数"),
            ("max_arrivals_per_replication", "100,000", "显式或估算的单次到达数"),
            ("max_estimated_events", "10,000,000", "所有场景、replication 的估算事件总量"),
        ],
        [3100, 1800, 4460],
        caption="表 16　PolicyLimits 默认值",
    )
    add_heading(doc, "事件估算", 2)
    add_body(
        doc,
        "若配置了 max_arrivals，直接使用该值；否则按 until ÷ 平均到达间隔 + 1 向上取整。平均到达间隔分别取指数 mean、定值 value、均匀分布 (low+high)/2、三角分布 (low+mode+high)/3。总估算事件近似为 arrivals × (station_count+1) × replications，并在所有场景间求和。",
    )
    add_callout(
        doc,
        "估算而非承诺",
        "该公式是服务保护预算，不是精确的 SimPy 事件数或性能 SLA。生产环境应按机器规格、模型复杂度与请求并发重新标定默认上限。",
        kind="info",
    )
    add_callout(
        doc,
        "作用范围",
        "PolicyLimits 由控制服务、REST 会话、MCP validate_project 与批准动作应用强制执行；普通 CLI 的 simlab validate / run 只执行配置模型校验，不自动套用这些服务端上限。",
        kind="warn",
    )


def add_workflow(doc: Document) -> None:
    add_heading(doc, "人工审批状态机、一致性与审计", 1, page_break=True)
    add_heading(doc, "提案状态", 2)
    add_table(
        doc,
        ["状态", "含义", "允许后续"],
        [
            ("pending", "已创建，等待人工决定", "approve 或 reject"),
            ("approved", "人工已批准，等待/正在本地应用", "apply"),
            ("rejected", "人工拒绝，终态", "无"),
            ("applied", "动作通过本地校验并写入下一版配置，终态", "无"),
            ("failed", "批准后应用失败并记录错误，终态", "基于当前配置重新提案"),
        ],
        [1500, 4900, 2960],
        caption="表 17　ProposalStatus 状态机",
    )
    add_heading(doc, "乐观版本", 2)
    add_body(
        doc,
        "会话 version 随每次提案创建、决定和应用单调增加。除创建会话外，控制面的写请求都携带 expected_version；如果与当前 version 不同，服务返回冲突而不是覆盖并发修改。客户端应重新 GET 会话、读取最新 workflow_version 与 config_hash，再决定是否重试。",
    )
    add_heading(doc, "幂等 operation_id", 2)
    add_body(
        doc,
        "除创建会话外，每个写操作要求调用方提供唯一 operation_id。相同 ID 与完全相同的操作指纹会重放原结果；相同 ID 被用于不同请求则产生 IdempotencyConflict。唯一性作用域是单个会话，而不是全局。服务层为工作流内部操作使用独立命名空间，避免“批准并应用”过程中的 ID 碰撞。",
    )
    add_body(
        doc,
        "典型版本变化为：创建会话得到 workflow_version=0；创建一项提案加 1；approve 接口内部先批准再应用，因此成功时再加 2；reject 加 1；run 只记录所用 workflow_version，本身不改变工作流版本。",
    )
    add_body(
        doc,
        "REST 的批准端点把批准与应用捆绑为一次服务操作，因此成功响应直接返回 applied；approved 状态主要出现在底层 workflow 与审计序列中。一个包含 N 个动作的 AI plan 在创建 proposal 时会使版本累计增加 N。",
    )
    add_heading(doc, "配置哈希与旧提案", 2)
    add_body(
        doc,
        "canonical_config_hash 对 ProjectConfig 的 JSON 表示按键排序、紧凑序列化，再计算 SHA-256。提案记录创建时哈希；批准时必须仍与当前哈希一致。即使 version 操作本身合法，旧配置证据下产生的提案也不能套用到新配置。",
    )
    add_heading(doc, "审计事件", 2)
    add_table(
        doc,
        ["事件", "触发时点", "常见 details"],
        [
            ("proposal_created", "提案建立", "status、action"),
            ("proposal_approved", "人工批准", "status、reason"),
            ("proposal_rejected", "人工拒绝", "status、reason"),
            ("proposal_applied", "配置应用成功", "status"),
            ("proposal_failed", "应用失败", "status、error"),
        ],
        [2600, 2800, 3960],
        caption="表 18　审批工作流审计事件",
    )
    add_body(
        doc,
        "每条审计记录包含 sequence、version、UTC timestamp、operation_id、proposal_id、actor 与 details。服务事件还覆盖会话创建、运行和 AI 计划等活动，可通过 REST audit 或 WebSocket 按 sequence 重放。",
    )


def add_api(doc: Document) -> None:
    add_heading(doc, "REST API 与 WebSocket 控制面", 1)
    add_body(
        doc,
        "FastAPI 控制面默认绑定 127.0.0.1:8000。设置 SIMLAB_API_TOKEN 后，受保护 HTTP 路由要求 Authorization: Bearer <token>；WebSocket 同样使用 Authorization 请求头，Token 不放在 URL 查询参数中。",
    )
    add_heading(doc, "端点目录", 2)
    add_table(
        doc,
        ["方法", "路径", "用途"],
        [
            ("GET", "/health", "进程健康检查"),
            ("POST", "/v1/sessions", "用 ProjectConfig 创建会话"),
            ("GET", "/v1/sessions/{session_id}", "读取版本、配置、提案与最近运行"),
            ("GET", "/v1/sessions/{session_id}/allowed-actions", "读取当前精确动作白名单"),
            ("POST", "/v1/sessions/{session_id}/proposals", "提交一项人工/可信系统提案"),
            (
                "POST",
                "/v1/sessions/{session_id}/proposals:generate",
                "基于最新 KPI 生成 OpenAI 动作计划",
            ),
            (
                "POST",
                "/v1/sessions/{session_id}/proposals/{proposal_id}:approve",
                "人工批准并尝试应用",
            ),
            ("POST", "/v1/sessions/{session_id}/proposals/{proposal_id}:reject", "人工拒绝"),
            ("POST", "/v1/sessions/{session_id}/runs", "以当前快照执行实验"),
            ("GET", "/v1/sessions/{session_id}/runs/{run_id}", "读取运行状态与输出位置"),
            ("GET", "/v1/sessions/{session_id}/audit", "读取服务事件与审批审计"),
            ("WS", "/v1/sessions/{session_id}/events?after_sequence=0", "重放并订阅有序事件"),
        ],
        [1100, 5190, 3070],
        caption="表 19　控制面端点",
        font_size=7.25,
    )
    add_heading(doc, "会话创建以外的写请求字段", 2)
    add_table(
        doc,
        ["字段", "作用", "规则"],
        [
            (
                "operation_id",
                "幂等键",
                "1–90 个受支持 ASCII 字符；simlab: 前缀保留；同一语义重试时复用",
            ),
            ("expected_version", "并发控制", "必须等于会话当前 workflow_version"),
            ("action", "人工提案数据", "仅 proposals 端点；严格判别联合"),
            (
                "objective / reason",
                "AI 目标或人工决定理由",
                "分别最多 4,000 / 2,000 字符，不是代码入口",
            ),
            ("workers", "运行并行度", "仅 runs 端点；1–16，默认 1"),
        ],
        [2000, 2600, 4760],
        caption="表 20　写请求的控制字段",
    )
    add_body(
        doc,
        "例外：POST /v1/sessions 的请求体只有 config 与可选 session_id，不要求 operation_id 或 expected_version。actor 不由请求体提交，而是由服务端认证上下文派生：无 Token 的本地模式为 local-user，Bearer 模式为 token-admin，并在 AI 路径追加 :openai。",
    )
    add_heading(doc, "典型 HTTP 状态", 2)
    add_table(
        doc,
        ["状态", "含义", "处理建议"],
        [
            (
                "200 / 201",
                "读取/创建或写操作成功",
                "保存返回的 workflow_version、id 与 config_hash",
            ),
            ("400", "配置、动作或工作量校验失败", "修正请求，不要盲目重试"),
            ("401", "Bearer Token 缺失或错误", "核对 SIMLAB_API_TOKEN 与请求头"),
            ("404", "会话、提案或运行不存在", "核对 ID 与服务实例"),
            ("409", "版本、幂等或状态转换冲突", "重新读取会话后再决定"),
            ("500", "仿真运行执行失败", "检查服务日志与 run.error"),
            ("502", "OpenAI 提案生成失败", "检查 key、额度、网络、模型或输出状态"),
            ("422", "FastAPI 请求模型校验失败", "修正字段类型、必填项或长度后重发"),
        ],
        [1300, 3300, 4760],
        caption="表 21　API 错误映射",
    )
    add_heading(doc, "创建会话与运行示例", 2)
    add_code_block(
        doc,
        r"""$headers = @{ Authorization = "Bearer $env:SIMLAB_API_TOKEN" }
$config = Get-Content .\examples\api_service_center.json -Raw | ConvertFrom-Json
$session = Invoke-RestMethod -Method Post -Uri http://127.0.0.1:8000/v1/sessions `
  -Headers $headers -ContentType 'application/json' `
  -Body (@{ session_id='demo-session'; config=$config } | ConvertTo-Json -Depth 20)

$run = Invoke-RestMethod -Method Post `
  -Uri "http://127.0.0.1:8000/v1/sessions/$($session.session_id)/runs" `
  -Headers $headers -ContentType 'application/json' `
  -Body (@{ operation_id='run-001'; expected_version=$session.workflow_version; workers=1 } | ConvertTo-Json)""",
        font_size=7.15,
    )
    add_callout(
        doc,
        "输出路径隔离",
        "控制 API 无视请求配置中的 experiment.output_dir。每次运行写入 SIMLAB_API_OUTPUT_ROOT / session_id / run_id，避免并发覆盖和利用路径字段穿越服务端目录。",
        kind="safe",
    )
    add_body(
        doc,
        "run 与 proposals:generate 在服务器线程中执行，避免阻塞 FastAPI 事件循环；但发起 HTTP 请求的客户端仍会等待该操作完成，并非提交后立即返回的后台作业。WebSocket 每 0.25 秒检查事件，约连续 15 秒无事件时发送 heartbeat。proposals:generate 只看 run_ids 中最后一次运行；若最后一次失败，传给 AI 的 KPI 为空，不会自动回退到更早的成功结果。",
    )


def add_mcp(doc: Document) -> None:
    add_heading(doc, "受限 MCP 接口", 1)
    add_body(
        doc,
        "simlab-mcp 是独立 stdio 进程，通过 SIMLAB_API_URL 和 SIMLAB_API_TOKEN 连接正在运行的 REST 服务。因此 MCP 与 REST 共享同一会话、版本、配置哈希和审计记录，而不是维护另一份隐藏状态。",
    )
    add_table(
        doc,
        ["MCP 工具", "能力", "权限边界"],
        [
            ("validate_project", "本地校验 ProjectConfig 与工作量", "不创建会话，不运行"),
            ("get_session", "读取指定 REST 会话", "只读"),
            ("get_allowed_actions", "读取当前动作目录", "只读"),
            ("submit_proposal", "提交一项 pending 提案", "仍需 expected_version 与人工审批"),
            ("list_events", "读取 after_sequence 之后的事件", "只读审计/进度"),
        ],
        [2600, 3100, 3660],
        caption="表 22　MCP v2 暴露工具",
    )
    add_callout(
        doc,
        "刻意不暴露",
        "MCP 没有 create_session、run、approve 或 reject。会话创建、仿真启动与人工决定必须经 REST/OpenAPI 完成，智能客户端不能用新建会话或自批提案绕过控制边界。",
        kind="risk",
    )
    add_code_block(
        doc,
        '$env:SIMLAB_API_URL = "http://127.0.0.1:8000"\n$env:SIMLAB_API_TOKEN = "与 REST 服务相同的管理 Token"\n.\\.venv\\Scripts\\simlab-mcp.exe',
    )


def add_operations(doc: Document) -> None:
    add_heading(doc, "Windows PowerShell：安装、运行与联调", 1)
    add_callout(
        doc,
        "命令起点",
        r"所有下列相对路径命令都从 D:\文件\SimPy\simpy-kpi-lab\simpy-kpi-lab 开始。路径含中文时使用 Set-Location -LiteralPath 和引号。",
        kind="info",
    )
    add_heading(doc, "首次安装", 2)
    add_code_block(
        doc,
        r"""Set-Location -LiteralPath 'D:\文件\SimPy\simpy-kpi-lab\simpy-kpi-lab'
py -3 -m venv .venv
.\.venv\Scripts\python.exe -m pip install -e ".[dev]"
.\.venv\Scripts\python.exe -m pytest
.\.venv\Scripts\ruff.exe check .""",
    )
    add_heading(doc, "只运行仿真与 KPI", 2)
    add_code_block(
        doc,
        r""".\.venv\Scripts\simlab.exe validate .\examples\service_center.yaml
.\.venv\Scripts\simlab.exe run .\examples\service_center.yaml --workers 1
# 仿真 + KPI 不需要 OPENAI_API_KEY""",
    )
    add_heading(doc, "调用 OpenAI 分析", 2)
    add_code_block(
        doc,
        r"""$secureKey = Read-Host '粘贴 OpenAI API Key' -AsSecureString
$env:OPENAI_API_KEY = [System.Net.NetworkCredential]::new('', $secureKey).Password
.\.venv\Scripts\simlab.exe run .\examples\service_center.yaml --analyze
# 或对已存在结果分析：
.\.venv\Scripts\simlab.exe analyze .\outputs\service_center\results.json `
  --question '比较服务水平、P95 周期时间和利用率，并指出权衡。'
Remove-Item Env:OPENAI_API_KEY -ErrorAction SilentlyContinue""",
    )
    add_heading(doc, "启动控制 API", 2)
    add_code_block(
        doc,
        r"""$env:SIMLAB_API_TOKEN = '请替换为随机长字符串'
$env:SIMLAB_API_OUTPUT_ROOT = 'outputs/api_sessions'
.\.venv\Scripts\simlab.exe serve --host 127.0.0.1 --port 8000
# 浏览器打开 http://127.0.0.1:8000/docs""",
    )
    add_heading(doc, "完整人工审批演示", 2)
    add_body(
        doc,
        "保持 API 窗口运行，在第二个 PowerShell 窗口进入同一项目根目录并执行：",
    )
    add_code_block(doc, r".\examples\human_in_the_loop.ps1")
    add_body(
        doc,
        "演示脚本会创建会话、运行实验、读取白名单、提交提案、执行人工批准/拒绝并查看审计。若同时配置 OPENAI_API_KEY，还可调用 proposals:generate；没有 key 时可使用手工结构化提案完成闭环。",
    )


def add_outputs(doc: Document) -> None:
    add_heading(doc, "输出文件与结果解读", 1)
    add_table(
        doc,
        ["文件", "粒度", "主要内容", "推荐用途"],
        [
            (
                "results.json",
                "整个实验",
                "配置、随机流、catalog、replications、summary",
                "完整审计与程序读取",
            ),
            (
                "replications.csv",
                "场景 × replication",
                "seed、参数与宽表 KPI",
                "诊断离群值、配对差值",
            ),
            (
                "summary.csv",
                "场景 × KPI",
                "角色、方向、n、均值、std、SE、CI、极值",
                "方案比较与报表",
            ),
            ("ai_analysis.json", "一次 AI 分析", "结构化发现、建议、风险", "机器消费"),
            ("ai_analysis.md", "一次 AI 分析", "可读摘要", "人工阅读"),
        ],
        [1900, 1900, 3500, 2060],
        caption="表 23　CLI 实验输出",
        font_size=7.7,
    )
    add_heading(doc, "推荐的比较顺序", 2)
    add_numbered(
        doc,
        [
            "先检查 n_missing、cycle_completion_fraction 与 censored_cycle_count，确认数据质量可接受。",
            "再看 primary KPI：throughput、p95_cycle_time、service_level。",
            "用 guardrail 检查 wip_end 与各工位 p95_wait_time 是否恶化。",
            "用 utilization、avg_queue_length、avg_wait_time 等 driver 定位瓶颈。",
            "比较置信区间与 replication 原始差值；不要只按均值高低排序。",
            "最后让 AI 总结权衡或提出下一轮有限动作；业务责任人仍需审核。",
        ],
    )
    add_heading(doc, "结果 JSON 的可复现字段", 2)
    add_bullets(
        doc,
        [
            "schema_version 当前为 1.1；消费者应检查版本再解析。",
            "generated_at 使用 UTC ISO 时间。",
            "config 保存实际执行配置，而不是只保存输入文件路径。",
            "random_streams.method 为 blake2b_namespaced_v1，并记录 common_random_numbers。",
            "每个 replication 保存 scenario、parameters、replication、seed 与 metrics。",
        ],
    )
    add_callout(
        doc,
        "归档提醒",
        "CLI 对同一 output_dir 重复运行会覆盖固定文件名；当前 results.json 也不记录包版本、Git commit 或源码哈希。正式实验应把输出目录、配置源文件、软件版本和提交哈希一起归档。API 运行目录虽按 run_id 隔离，结果内的 experiment.output_dir 仍是输入值，实际路径以 RunRecord.output_dir 为准。",
        kind="warn",
    )


def add_troubleshooting(doc: Document) -> None:
    add_heading(doc, "故障排查", 1)
    add_table(
        doc,
        ["现象", "最常见原因", "处理"],
        [
            (
                "OpenAI API 认证失败",
                "变量名写错、key 无效/撤销、复制了空格",
                "检查 $env:OPENAI_API_KEY；重新创建 key；不要写反斜杠",
            ),
            (
                "未设置 OPENAI_API_KEY",
                "执行了 --analyze 或 proposals:generate",
                "仅仿真时去掉 AI 参数；需要时再临时设置 key",
            ),
            (
                "HTTP 409 version conflict",
                "expected_version 已过期",
                "GET 最新会话，重新评估操作，不原样盲重试",
            ),
            (
                "HTTP 409 idempotency conflict",
                "operation_id 被不同请求复用",
                "为新语义生成新 ID；网络重试才复用旧 ID",
            ),
            ("stale proposal", "配置哈希在提案后已变化", "基于当前 config/KPI 重新生成提案"),
            (
                "HTTP 400 workload",
                "场景、replication、到达或事件估算超限",
                "缩小网格/时长/次数，或经容量评估后调整服务端限制",
            ),
            (
                "MCP 无法连接",
                "REST 未启动、URL/Token 不一致",
                "检查 SIMLAB_API_URL、SIMLAB_API_TOKEN 与 /health",
            ),
            (
                "WebSocket 4401 / 4404",
                "认证失败 / 会话不存在",
                "使用 Authorization 头并核对 session_id",
            ),
            (
                "中文控制台乱码",
                "终端编码不一致",
                "PowerShell 7 优先；必要时执行 chcp 65001 并统一 UTF-8",
            ),
        ],
        [2400, 3200, 3760],
        caption="表 24　常见问题与处理",
        font_size=7.25,
    )
    add_heading(doc, "认证失败的最短核查", 2)
    add_code_block(
        doc,
        r"""# 只检查变量是否存在，不要把 key 输出到屏幕
if ([string]::IsNullOrWhiteSpace($env:OPENAI_API_KEY)) {
  'OPENAI_API_KEY 未设置'
} else {
  "OPENAI_API_KEY 已设置，长度=$($env:OPENAI_API_KEY.Length)"
}
# 当前窗口临时移除
Remove-Item Env:OPENAI_API_KEY -ErrorAction SilentlyContinue""",
    )
    add_callout(
        doc,
        "密钥获取",
        "API Key 在 OpenAI Platform 的 API Keys 页面创建，并受对应 API 项目的权限与用量设置约束。创建后应立即存入安全的密码或密钥管理工具；仅登录 ChatGPT 不会自动为 PowerShell 设置 OPENAI_API_KEY。",
        kind="info",
    )


def add_validation(doc: Document) -> None:
    add_heading(doc, "测试覆盖、安全检查与当前限制", 1, page_break=True)
    add_heading(doc, "当前验证基线", 2)
    add_body(
        doc,
        "当前仓库已通过 72 项 pytest 测试，并通过 Ruff 检查、格式检查、pip 依赖一致性与安全回归扫描。这个数字是文档生成时的基线；以后新增代码后应以实际测试输出为准。",
    )
    add_table(
        doc,
        ["领域", "重点验证"],
        [
            ("配置", "严格字段、分布参数、warmup/until、唯一工位名、网格路径"),
            ("仿真/KPI", "可重复 seed、窗口边界、队列面积、利用率、右删失、空样本"),
            ("实验", "场景笛卡尔积、CRN、并行、聚合缺失值、JSON/CSV"),
            ("OpenAI", "Structured Outputs、认证/限流/超时/状态、越权 target、审批标志"),
            ("工作流", "状态转换、version、幂等重放/冲突、审计序列"),
            ("策略", "allowlist、严格应用、哈希、网格工作量、资源重排"),
            ("API/MCP", "认证、错误映射、事件流、输出隔离、受限工具目录"),
            ("安全回归", "无 eval/exec/shell 执行入口，无真实密钥入库"),
        ],
        [2500, 6860],
        caption="表 25　测试与安全检查范围",
    )
    add_heading(doc, "发布前命令", 2)
    add_code_block(
        doc,
        r""".\.venv\Scripts\python.exe -m pytest
.\.venv\Scripts\ruff.exe check .
.\.venv\Scripts\ruff.exe format --check .
.\.venv\Scripts\python.exe -m pip check""",
    )
    add_heading(doc, "当前限制", 2)
    add_bullets(
        doc,
        [
            "会话、幂等记录、提案与事件保存在单进程内存；服务重启后丢失。",
            "WebSocket broker 面向单个 Uvicorn worker；多实例间不会自动共享事件。",
            "认证是单一管理员 Bearer Token，尚无 approver / runner / viewer 角色分离。",
            "仿真不支持暂停、恢复、快照或运行中热改参数。",
            "业务模型是串行同一路由，不含条件路由、优先级、故障、批处理或资源班次。",
            "统计置信区间为正态近似，尚未内置 t 区间、bootstrap、多重比较或自动停止规则。",
        ],
    )


def add_extension(doc: Document) -> None:
    add_heading(doc, "安全扩展指南与生产路线", 1)
    add_heading(doc, "扩展业务模型", 2)
    add_numbered(
        doc,
        [
            '先在 config.py 添加明确类型和约束，继续使用 extra="forbid"。',
            "在 simulation.py 中以事件和过程表达路由、返工或班次，不把业务逻辑放进 CLI。",
            "为每个新增随机过程分配稳定命名空间，避免扰动已有随机序列。",
            "在 KPICollector 的事件边界采集原始事实，再由 finalize 计算指标。",
            "为新增 KPI 扩展 metric catalog，给出 role、direction、unit 与 definition。",
            "若允许 AI 修改新参数，逐项加入动态 allowlist，并补充负向越权测试与工作量预算。",
        ],
    )
    add_heading(doc, "新增动作类型时的安全门", 2)
    add_bullets(
        doc,
        [
            "动作必须是可序列化、有限字段的 Pydantic 判别联合，不接受代码字符串。",
            "先定义精确 target 目录，再写纯本地 apply 函数；禁止动态 import、反射执行或任意 dotted path。",
            "把动作应用到配置副本，完整重验后才替换当前状态。",
            "定义人类可理解的风险、预期效果、回滚方式与审计 details。",
            "补充版本冲突、幂等冲突、旧哈希、非法类型、极端数值和顺序依赖测试。",
        ],
    )
    add_heading(doc, "生产化优先级", 2)
    add_table(
        doc,
        ["优先级", "建设项", "目的"],
        [
            ("P0", "PostgreSQL/SQLite 持久化会话、操作与 outbox 事件", "重启可恢复、事务后广播"),
            ("P0", "TLS、反向代理、正式身份认证与 RBAC", "区分提案、运行、审批、查看权限"),
            ("P1", "Redis/消息总线与后台任务队列", "多实例 WebSocket 与长任务可靠执行"),
            ("P1", "Secrets Manager 与密钥轮换", "避免环境变量泄露与人工复制"),
            ("P1", "限流、并发配额、超时取消与指标监控", "保护服务容量与成本"),
            ("P2", "t/bootstrap CI、配对差值、功效分析与自动停止", "提高实验结论质量"),
            ("P2", "配置/结果 schema 迁移与模型注册表", "长期兼容与审计"),
        ],
        [1200, 5000, 3160],
        caption="表 26　建议的生产化路线",
        font_size=7.65,
    )


def add_appendices(doc: Document) -> None:
    add_heading(doc, "附录：项目文件地图", 1)
    add_table(
        doc,
        ["路径", "职责"],
        [
            ("src/simlab/config.py", "配置模型、分布采样与 YAML 读取"),
            ("src/simlab/rng.py", "BLAKE2b 稳定 seed 派生"),
            ("src/simlab/simulation.py", "单个 SimPy replication"),
            ("src/simlab/kpi.py", "KPICollector 与 metric catalog"),
            ("src/simlab/experiment.py", "场景、replication、聚合、导出"),
            ("src/simlab/ai.py", "KPI 结构化分析"),
            ("src/simlab/agent.py", "OpenAI 动作提案"),
            ("src/simlab/policy.py", "allowlist、配置哈希、应用与工作量限制"),
            ("src/simlab/workflow.py", "人工审批状态机、版本、幂等、审计"),
            ("src/simlab/service.py", "会话、运行、AI 计划与服务事件"),
            ("src/simlab/api.py", "FastAPI REST 与 WebSocket"),
            ("src/simlab/mcp_server.py", "受限 MCP v2 客户端"),
            ("src/simlab/cli.py", "validate / run / analyze / serve 命令"),
            ("examples/api_service_center.json", "控制 API 示例项目配置"),
            ("examples/human_in_the_loop.ps1", "PowerShell 人工审批演示"),
            ("tests/", "功能、并发、安全与接口回归测试"),
        ],
        [4300, 5060],
        caption="表 27　当前仓库主要文件",
        font_size=7.6,
    )

    add_heading(doc, "附录：验收清单", 1, page_break=True)
    add_table(
        doc,
        ["检查项", "通过标准"],
        [
            ("本地仿真", "无 key 时 validate/run 成功并产生三类结果文件"),
            ("复现", "同配置、同版本、同 seed 的 replication KPI 一致"),
            ("多场景", "网格场景数与笛卡尔积一致，CRN 设置符合实验意图"),
            ("KPI 质量", "n_missing、删失比例和窗口口径被检查"),
            ("OpenAI", "真实 key 仅在环境变量；结构化响应通过 Pydantic"),
            ("安全提案", "越权 target、requires_approval=false、NaN/Infinity 均拒绝"),
            ("审批", "旧 version、旧 config_hash、ID 复用冲突均不会覆盖状态"),
            ("API", "Token、错误映射、输出隔离与事件重放通过"),
            ("MCP", "工具目录不包含 create/run/approve/reject"),
            ("发布", "pytest、Ruff、pip check 与密钥扫描通过"),
        ],
        [3300, 6060],
        caption="表 28　功能与安全验收清单",
    )

    add_heading(doc, "附录：术语表", 1)
    add_table(
        doc,
        ["术语", "解释"],
        [
            ("replication", "同一场景下使用独立随机种子完成的一次完整仿真"),
            ("scenario", "parameter_grid 中一组具体参数组合"),
            ("warmup", "为降低空系统初始偏差而排除的早期统计时间段"),
            ("cohort", "按到达与完成条件定义、用于周期指标的一组实体"),
            ("右删失", "until 时实体尚未完成，完整周期时间不可观测"),
            ("CRN", "Common Random Numbers；让方案共享随机输入以提高差值比较精度"),
            ("Structured Outputs", "让模型输出遵循 JSON Schema，并由 Pydantic 解析"),
            ("allowlist", "本地代码生成的精确 action_type + target 允许集合"),
            ("operation_id", "会话创建以外写操作的幂等标识；REST 限 1–90 字符"),
            ("config_hash", "规范化配置 JSON 的 SHA-256，用于识别提案是否过期"),
        ],
        [2500, 6860],
        caption="表 29　关键术语",
    )

    add_heading(doc, "附录：参考资料", 1)
    add_body(
        doc, "项目内资料：README.md、pyproject.toml、src/simlab/、examples/ 与 tests/。外部资料："
    )
    sources = [
        (
            "OpenAI Responses API 参考",
            "https://developers.openai.com/api/reference/python/resources/responses",
        ),
        (
            "OpenAI Structured Outputs 指南",
            "https://developers.openai.com/api/docs/guides/structured-outputs",
        ),
        ("OpenAI API Keys", "https://platform.openai.com/api-keys"),
        ("SimPy 官方文档", "https://simpy.readthedocs.io/"),
        ("FastAPI 官方文档", "https://fastapi.tiangolo.com/"),
        ("Model Context Protocol Python SDK", "https://github.com/modelcontextprotocol/python-sdk"),
    ]
    for label, url in sources:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, _BULLET_NUM_ID or 0)
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)
        set_paragraph_language(paragraph)
        add_hyperlink(paragraph, label, url)

    add_callout(
        doc,
        "文档基线",
        f"本说明依据 {DOCUMENT_DATE.isoformat()} 工作区中的 v{DOCUMENT_VERSION} 源代码生成。若代码或依赖版本变化，应重新运行测试并同步修订本文中的接口、限制和测试数量。",
        kind="info",
    )


def build_document() -> Document:
    global _BULLET_NUM_ID, _DECIMAL_NUM_ID

    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    set_update_fields_on_open(doc)
    _BULLET_NUM_ID = create_numbering_definition(doc, bullet=True)
    _DECIMAL_NUM_ID = create_numbering_definition(doc, bullet=False)

    properties = doc.core_properties
    properties.title = "SimPy KPI Lab 安全闭环模型详细解释说明"
    properties.subject = "SimPy、KPI、多次实验、OpenAI 结构化提案、人工审批、REST/WebSocket 与 MCP"
    properties.author = "SimPy KPI Lab"
    properties.keywords = "SimPy, KPI, OpenAI, Structured Outputs, HITL, FastAPI, MCP"
    properties.comments = f"Design preset: {DESIGN_PRESET}; header pattern: {HEADER_PATTERN}."

    add_cover(doc)
    add_reader_guide(doc)
    add_positioning(doc)
    add_architecture(doc)
    add_configuration(doc)
    add_simulation(doc)
    add_experiments(doc)
    add_kpis(doc)
    add_openai(doc)
    add_policy(doc)
    add_workload(doc)
    add_workflow(doc)
    add_api(doc)
    add_mcp(doc)
    add_operations(doc)
    add_outputs(doc)
    add_troubleshooting(doc)
    add_validation(doc)
    add_extension(doc)
    add_appendices(doc)
    return doc


def main() -> None:
    QA_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
